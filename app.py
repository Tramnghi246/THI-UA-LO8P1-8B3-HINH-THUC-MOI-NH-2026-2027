import io
import os
import time
import pandas as pd
import requests
import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# 1. CẤU HÌNH TRANG MOBILE GỌN GÀNG (KHÔNG SIDEBAR, BẤM LÀ CHẠY NGAY)
# ==============================================================================
st.set_page_config(
    page_title="Lớp 8B3 Vươn Mình Thế Giới",
    page_icon="🌟",
    layout="centered",
    initial_sidebar_state="collapsed",
)

API_URL = "https://script.google.com/macros/s/AKfycbwoy6GT8XEkwwrLTSXrIk3tMj6lzATeiwrw02cWYz7yEXDhzEtDOoyL8VssgXuJE08q/exec"
DATA_FILE = "du_lieu_thi_dua.csv"

# ==============================================================================
# 2. BỘ CSS CHUẨN GIAO DIỆN APP MOBILE
# ==============================================================================
st.markdown(
    """
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Segoe+UI:wght@400;600;700;800&display=swap');
        html, body, [class*="css"] {
            font-family: 'Segoe UI', system-ui, -apple-system, sans-serif;
        }

        /* Ẩn hoàn toàn sidebar */
        section[data-testid="stSidebar"], [data-testid="collapsedControl"] {
            display: none !important;
        }

        .block-container {
            padding-top: 1rem !important;
            padding-bottom: 5.5rem !important;
            max-width: 600px !important;
        }

        /* Header trên cùng */
        .app-topbar-m {
            background: linear-gradient(135deg, #0b4f8a 0%, #1976d2 100%);
            color: #ffffff;
            text-align: center;
            padding: 18px 14px;
            border-radius: 16px;
            margin-bottom: 14px;
            box-shadow: 0 4px 14px rgba(11, 79, 138, 0.25);
        }
        .app-title-m {
            font-size: 1.25rem;
            font-weight: 800;
            letter-spacing: 0.5px;
            text-transform: uppercase;
        }
        .app-sub-m {
            font-size: 0.8rem;
            color: #fde047;
            font-weight: 600;
            margin-top: 2px;
        }

        /* Thẻ học sinh */
        .student-m-card {
            background: #ffffff;
            border-radius: 14px;
            padding: 12px 16px;
            margin-bottom: 10px;
            border: 1px solid #e2e8f0;
            box-shadow: 0 2px 6px rgba(0,0,0,0.03);
            display: flex;
            justify-content: space-between;
            align-items: center;
        }
        .student-m-avatar {
            width: 42px;
            height: 42px;
            background: #e0f2fe;
            color: #0284c7;
            border-radius: 50%;
            display: flex;
            align-items: center;
            justify-content: center;
            font-weight: 800;
            font-size: 1.15rem;
            margin-right: 12px;
        }

        /* Thẻ thống kê */
        .stat-grid-m {
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 10px;
            margin-bottom: 16px;
        }
        .stat-box-m {
            background: #ffffff;
            border-radius: 14px;
            padding: 14px;
            border: 1px solid #e2e8f0;
            text-align: center;
            box-shadow: 0 2px 6px rgba(0,0,0,0.03);
        }
        .stat-box-label {
            font-size: 0.75rem;
            color: #64748b;
            font-weight: 700;
            text-transform: uppercase;
        }
        .stat-box-val {
            font-size: 1.6rem;
            font-weight: 800;
            color: #0f172a;
            margin-top: 2px;
        }

        /* Nút bấm */
        .stButton>button {
            border-radius: 10px;
            font-weight: 700;
            padding: 8px 12px;
            transition: all 0.2s ease;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

# ==============================================================================
# 3. DANH MỤC THI ĐUA & DỮ LIỆU KHỞI TẠO
# ==============================================================================
DANH_SACH_LOI = {
    "🌸 Thầy cô tuyên dương (+2)": 2,
    "🌟 Phát biểu xây dựng bài (+1)": 1,
    "💯 Đạt điểm 9, 10 (+2)": 2,
    "🏆 Đạt giải phong trào (+5)": 5,
    "🤬 Nói tục (-5)": -5,
    "📱 Làm việc riêng (-5)": -5,
    "🔊 Nói chuyện trong giờ (-2)": -2,
    "✉️ Nghỉ học có phép (-2)": -2,
    "⏰ Đi muộn (-1)": -1,
    "📚 Không học bài / Thiếu BTVN (-5)": -5,
    "👔 Không mặc đồng phục (-5)": -5,
    "🧹 Không trực nhật / vệ sinh (-5)": -5,
    "📱 Sử dụng điện thoại không xin phép (-15)": -15,
    "🚨 Nghỉ học không phép (-5)": -5,
}

DEFAULT_STUDENTS = [
    {"STT": 1, "Họ và tên": "Huỳnh Bảo Ngọc Thiên Ân", "Tổ": "Tổ 1", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 2, "Họ và tên": "Lê Minh Thái Dương", "Tổ": "Tổ 1", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 3, "Họ và tên": "Ngô Gia Hân", "Tổ": "Tổ 1", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 4, "Họ và tên": "Phạm Hoàng Đăng Khôi", "Tổ": "Tổ 1", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 5, "Họ và tên": "Trần Thị Cẩm Ly", "Tổ": "Tổ 1", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 6, "Họ và tên": "Vũ Đức Trọng", "Tổ": "Tổ 1", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 7, "Họ và tên": "Nguyễn Tuyền An", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 8, "Họ và tên": "Nguyễn Thị Ngọc Hiên", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 9, "Họ và tên": "Huỳnh Ngọc Huy", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 10, "Họ và tên": "Nguyễn Huỳnh Gia Hưng", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 11, "Họ và tên": "Trần Quốc Bảo", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 12, "Họ và tên": "Nguyễn Minh Khang", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 13, "Họ và tên": "Bùi Thị Tuyết Mai", "Tổ": "Tổ 2", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 14, "Họ và tên": "Lê Tuấn Hưng", "Tổ": "Tổ 3", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 15, "Họ và tên": "Lê Quỳnh Diễm Hương", "Tổ": "Tổ 3", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 16, "Họ và tên": "Nguyễn Lê Hoàng Mai", "Tổ": "Tổ 3", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 17, "Họ và tên": "Đỗ Hoàng Nam", "Tổ": "Tổ 3", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 18, "Họ và tên": "Võ Thị Yến Nhi", "Tổ": "Tổ 3", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 19, "Họ và tên": "Phan Thanh Phong", "Tổ": "Tổ 3", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 20, "Họ và tên": "Võ Huỳnh Minh An", "Tổ": "Tổ 4", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 21, "Họ và tên": "Lê Văn Tuấn", "Tổ": "Tổ 4", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 22, "Họ và tên": "Đặng Ngọc Như Quỳnh", "Tổ": "Tổ 4", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 23, "Họ và tên": "Nguyễn Tấn Tài", "Tổ": "Tổ 4", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 24, "Họ và tên": "Lê Thị Thảo Trang", "Tổ": "Tổ 4", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""},
    {"STT": 25, "Họ và tên": "Nguyễn Hoàng Việt", "Tổ": "Tổ 4", "Điểm thi đua": 100, "Xếp loại": "Tốt", "Lỗi vi phạm": ""}
]

# ==============================================================================
# 4. HÀM XỬ LÝ DỮ LIỆU
# ==============================================================================
def tinh_xep_loai(diem):
    try:
        diem = float(diem)
    except Exception:
        diem = 100.0
    if diem >= 100:
        return "Tốt"
    elif diem >= 85:
        return "Khá"
    elif diem >= 70:
        return "Trung bình"
    else:
        return "Yếu"

def clean_dataframe(df_input):
    if df_input is None or df_input.empty:
        return pd.DataFrame(DEFAULT_STUDENTS)

    rename_dict = {}
    for col in df_input.columns:
        c_low = str(col).strip().lower()
        if c_low in ["họ và tên", "họ tên", "tên học sinh", "ten"]:
            rename_dict[col] = "Họ và tên"
        elif c_low in ["tổ", "to", "nhóm"]:
            rename_dict[col] = "Tổ"
        elif c_low in ["điểm thi đua", "điểm", "diem", "diem thi dua"]:
            rename_dict[col] = "Điểm thi đua"
        elif c_low in ["xếp loại", "xep loai", "loại"]:
            rename_dict[col] = "Xếp loại"
        elif any(k in c_low for k in ["lỗi", "vi phạm", "ghi chú", "tuyên dương", "chi tiết", "nội dung"]):
            rename_dict[col] = "Lỗi vi phạm"

    if rename_dict:
        df_input = df_input.rename(columns=rename_dict)

    if "Lỗi vi phạm" not in df_input.columns:
        if len(df_input.columns) >= 6:
            df_input["Lỗi vi phạm"] = df_input.iloc[:, 5]
        else:
            df_input["Lỗi vi phạm"] = ""

    df_clean = df_input.dropna(subset=["Họ và tên"]).copy()
    df_clean["Họ và tên"] = df_clean["Họ và tên"].astype(str).str.strip()
    df_clean["Tổ"] = df_clean["Tổ"].fillna("Tổ 1").astype(str).str.strip()
    df_clean["Điểm thi đua"] = pd.to_numeric(df_clean["Điểm thi đua"], errors="coerce").fillna(100).astype(int)
    df_clean["Xếp loại"] = df_clean["Điểm thi đua"].apply(tinh_xep_loai)
    df_clean["Lỗi vi phạm"] = df_clean["Lỗi vi phạm"].astype(str).str.strip().replace(["nan", "None", "NaN"], "")
    df_clean = df_clean.reset_index(drop=True)
    df_clean["STT"] = list(range(1, len(df_clean) + 1))
    return df_clean

def load_data():
    if API_URL and "/exec" in API_URL:
        try:
            res = requests.get(API_URL, params={"_t": time.time()}, timeout=7)
            if res.status_code == 200:
                data = res.json()
                if isinstance(data, list) and len(data) > 1:
                    df_temp = pd.DataFrame(data[1:], columns=data[0])
                    df_temp = clean_dataframe(df_temp)
                    df_temp.to_csv(DATA_FILE, index=False)
                    return df_temp
        except Exception:
            pass

    if os.path.exists(DATA_FILE):
        try:
            return clean_dataframe(pd.read_csv(DATA_FILE))
        except Exception:
            pass

    return pd.DataFrame(DEFAULT_STUDENTS)

def update_gsheet(name, score, xeploai, log_text):
    if not API_URL or "/exec" not in API_URL:
        return
    try:
        params = {
            "action": "update",
            "name": name,
            "score": score,
            "xeploai": xeploai,
            "log": log_text,
            "_t": time.time(),
        }
        requests.get(API_URL, params=params, timeout=10)
    except Exception:
        pass

if "df" not in st.session_state:
    st.session_state.df = load_data()

df = st.session_state.df

# Khởi tạo tab hiện tại: mặc định là "Chỉnh sửa điểm các tổ"
if "mobile_tab" not in st.session_state:
    st.session_state.mobile_tab = "Chỉnh sửa điểm các tổ"

# Số lần vi phạm/thưởng mặc định trong phiên
if "record_times" not in st.session_state:
    st.session_state.record_times = 1

# ==============================================================================
# 5. HEADER MOBILE: LỚP 8B3 VƯƠN MÌNH THẾ GIỚI
# ==============================================================================
st.markdown(
    """
    <div class="app-topbar-m">
        <div class="app-title-m">🌟 LỚP 8B3 VƯƠN MÌNH THẾ GIỚI</div>
        <div class="app-sub-m">TRƯỜNG THCS CHÍ CÔNG • NĂM HỌC 2026 - 2027</div>
    </div>
    """,
    unsafe_allow_html=True,
)

# Thanh chọn vai trò nhanh & nút làm mới
role_c1, role_c2 = st.columns([2, 1])
with role_c1:
    vai_tro = st.selectbox(
        "Vai trò đang làm việc:",
        ["👑 Giáo viên chủ nhiệm", "Tổ trưởng Tổ 1", "Tổ trưởng Tổ 2", "Tổ trưởng Tổ 3", "Tổ trưởng Tổ 4"],
        label_visibility="collapsed",
    )
with role_c2:
    if st.button("🔄 Làm mới", use_container_width=True):
        st.session_state.df = load_data()
        st.toast("Đã cập nhật dữ liệu mới nhất!", icon="🔄")
        time.sleep(0.5)
        st.rerun()

is_gvcn = "Giáo viên" in vai_tro

# Lọc theo tổ
if "Tổ 1" in vai_tro:
    df_role = df[df["Tổ"].astype(str).str.contains("1", na=False)].copy()
elif "Tổ 2" in vai_tro:
    df_role = df[df["Tổ"].astype(str).str.contains("2", na=False)].copy()
elif "Tổ 3" in vai_tro:
    df_role = df[df["Tổ"].astype(str).str.contains("3", na=False)].copy()
elif "Tổ 4" in vai_tro:
    df_role = df[df["Tổ"].astype(str).str.contains("4", na=False)].copy()
else:
    df_role = df.copy()

st.markdown("<hr style='margin: 8px 0 14px 0; border: none; border-top: 1px solid #e2e8f0;'>", unsafe_allow_html=True)

# ==============================================================================
# 6. NỘI DUNG 3 TÍNH NĂNG CHÍNH DƯỚI ĐÁY
# ==============================================================================

# --- MÀN HÌNH 1: THỐNG KÊ ---
if st.session_state.mobile_tab == "Thống kê":
    st.markdown("##### 📊 Bảng Thống Kê Lớp 8B3")

    total_students = len(df)
    avg_score = round(df["Điểm thi đua"].mean(), 1) if not df.empty else 0
    tot_count = len(df[df["Xếp loại"] == "Tốt"])
    yeu_count = len(df[df["Xếp loại"].isin(["Trung bình", "Yếu"])])

    st.markdown(
        f"""
        <div class="stat-grid-m">
            <div class="stat-box-m">
                <div class="stat-box-label">Sĩ số lớp</div>
                <div class="stat-box-val">{total_students}</div>
            </div>
            <div class="stat-box-m">
                <div class="stat-box-label">Điểm TB Lớp</div>
                <div class="stat-box-val" style="color:#0284c7;">{avg_score}</div>
            </div>
            <div class="stat-box-m">
                <div class="stat-box-label">Xếp loại Tốt</div>
                <div class="stat-box-val" style="color:#16a34a;">{tot_count}</div>
            </div>
            <div class="stat-box-m">
                <div class="stat-box-label">Cần nhắc nhở</div>
                <div class="stat-box-val" style="color:#dc2626;">{yeu_count}</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    if "Tổ" in df.columns and not df.empty:
        st.markdown("###### 🏆 Điểm trung bình thi đua 4 Tổ:")
        avg_by_to = df.groupby("Tổ")["Điểm thi đua"].mean().round(1).reset_index()
        avg_by_to = avg_by_to.sort_values(by="Điểm thi đua", ascending=False)
        st.bar_chart(data=avg_by_to.set_index("Tổ"), use_container_width=True)

# --- MÀN HÌNH 2: XUẤT FILE TỔNG HỢP ---
elif st.session_state.mobile_tab == "Xuất file tổng hợp":
    st.markdown("##### 📥 Xuất Báo Cáo Xếp Hạng Thi Đua")
    st.caption("Danh sách tự động xếp hạng học sinh đạt điểm cao nhất từ trên xuống dưới.")

    def tao_excel_chuyen_nghiep(dataframe):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "BaoCao8B3"
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Calibri", size=15, bold=True, color="0B4F8A")
        font_sub = Font(name="Calibri", size=11, italic=True, color="475569")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_data = Font(name="Calibri", size=11)
        font_bold = Font(name="Calibri", size=11, bold=True)

        fill_header = PatternFill(start_color="0B4F8A", end_color="0B4F8A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        border_thin = Side(border_style="thin", color="CBD5E1")
        border_full = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        ws.merge_cells("A1:F1")
        ws["A1"] = "TRƯỜNG THCS CHÍ CÔNG - LỚP 8B3 VƯƠN MÌNH THẾ GIỚI"
        ws["A1"].font = font_title
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:F2")
        ws["A2"] = f"BẢNG TỔNG HỢP XẾP HẠNG THI ĐUA (Xuất ngày: {time.strftime('%d/%m/%Y %H:%M')})"
        ws["A2"].font = font_sub
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

        ws.append([])

        headers = ["Hạng", "Họ và tên", "Tổ", "Điểm thi đua", "Xếp loại", "Chi tiết vi phạm / Tuyên dương"]
        ws.append(headers)
        header_row = 4

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_full

        ws.row_dimensions[header_row].height = 26

        df_sorted = dataframe.sort_values(by=["Điểm thi đua", "Họ và tên"], ascending=[False, True]).reset_index(drop=True)

        current_row = 5
        for idx, row in df_sorted.iterrows():
            raw_val = str(row.get("Lỗi vi phạm", "")).strip()
            score_val = int(row.get("Điểm thi đua", 100))
            
            if raw_val in ["None", "nan", "NaN", ""]:
                if score_val > 100:
                    noi_dung_xuat = f"Tuyên dương (+{score_val - 100} điểm)"
                elif score_val < 100:
                    noi_dung_xuat = f"Vi phạm nề nếp ({score_val - 100} điểm)"
                else:
                    noi_dung_xuat = "—"
            else:
                noi_dung_xuat = raw_val

            ws.append([
                idx + 1,
                row.get("Họ và tên", ""),
                row.get("Tổ", ""),
                score_val,
                row.get("Xếp loại", "Tốt"),
                noi_dung_xuat
            ])

            for c_idx in range(1, 7):
                cell = ws.cell(row=current_row, column=c_idx)
                cell.font = font_data
                cell.border = border_full
                if current_row % 2 == 0:
                    cell.fill = fill_zebra

                if c_idx in [1, 3, 5]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx == 4:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 6:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws.row_dimensions[current_row].height = 24
            current_row += 1

        ws.merge_cells(f"A{current_row}:C{current_row}")
        total_label = ws.cell(row=current_row, column=1, value="ĐIỂM TRUNG BÌNH CẢ LỚP")
        total_label.font = font_bold
        total_label.alignment = Alignment(horizontal="center", vertical="center")

        avg_val = round(dataframe["Điểm thi đua"].mean(), 1) if not dataframe.empty else 0
        avg_cell = ws.cell(row=current_row, column=4, value=avg_val)
        avg_cell.font = font_bold
        avg_cell.alignment = Alignment(horizontal="right", vertical="center")

        for c_idx in range(1, 7):
            c = ws.cell(row=current_row, column=c_idx)
            c.fill = fill_total
            c.border = border_full

        ws.row_dimensions[current_row].height = 24

        col_widths = {1: 8, 2: 26, 3: 12, 4: 15, 5: 14, 6: 60}
        for col_idx, width in col_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    excel_data = tao_excel_chuyen_nghiep(df)

    st.download_button(
        label="📊 TẢI XUỐNG FILE EXCEL (.XLSX)",
        data=excel_data,
        file_name=f"Bao_Cao_Thi_Dua_8B3_{time.strftime('%Y_%m_%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )

# --- MÀN HÌNH 3: CHỈNH SỬA ĐIỂM CÁC TỔ (CÓ NÚT TĂNG GIẢM [+] [-]) ---
elif st.session_state.mobile_tab == "Chỉnh sửa điểm các tổ":
    st.markdown(f"##### ✏️ Chỉnh Sửa & Chấm Điểm ({vai_tro})")

    # Chọn học sinh
    danh_sach_hs = [
        f"#{row['STT']} - {row['Họ và tên']} ({row['Tổ']})"
        for _, row in df_role.iterrows()
    ]

    selected_student_raw = st.selectbox(
        "👤 Chọn học sinh cần chấm / sửa điểm:",
        danh_sach_hs if danh_sach_hs else ["(Không có học sinh trong phân ban này)"],
    )

    if danh_sach_hs and not selected_student_raw.startswith("("):
        stt_hs = int(selected_student_raw.split(" - ")[0].replace("#", ""))
        match_idx = df[df["STT"] == stt_hs].index[0]
        cur_student = df.loc[match_idx]

        curr_score = int(cur_student["Điểm thi đua"])
        curr_rank = cur_student["Xếp loại"]
        score_color = "#16a34a" if curr_score >= 100 else ("#d97706" if curr_score >= 85 else "#dc2626")

        # Thẻ thông tin học sinh
        st.markdown(
            f"""
            <div class="student-m-card">
                <div style="display: flex; align-items: center;">
                    <div class="student-m-avatar">{cur_student['Họ và tên'].split()[-1][0].upper()}</div>
                    <div>
                        <div style="font-weight: 700; color: #1e293b;">{cur_student['Họ và tên']}</div>
                        <div style="font-size: 0.8rem; color: #64748b;">{cur_student['Tổ']} • STT #{cur_student['STT']}</div>
                    </div>
                </div>
                <div style="text-align: right;">
                    <div style="font-size: 1.3rem; font-weight: 800; color: {score_color};">{curr_score} đ</div>
                    <span style="font-size: 0.75rem; font-weight: 700; color: {score_color};">Xếp loại: {curr_rank}</span>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        # 2 Tab con: [Ghi Vi Phạm / Thưởng] vs [Sửa Điểm Trực Tiếp]
        mode_action = st.radio("Thao tác:", ["🎯 Ghi vi phạm / Thưởng", "🛠️ Sửa điểm trực tiếp"], horizontal=True, label_visibility="collapsed")

        # KHỐI 1: GHI NHẬN CÓ BỘ NÚT TĂNG GIẢM [+] [-] SỐ LẦN
        if mode_action == "🎯 Ghi vi phạm / Thưởng":
            with st.container(border=True):
                loi_chon = st.selectbox("Nội dung:", list(DANH_SACH_LOI.keys()))
                
                # Cụm nút TĂNG GIẢM SỐ LẦN BẰNG NÚT BẤM [+] [-]
                st.markdown("<div style='text-align: center; font-weight: 700; font-size: 0.9rem;'>Số lần vi phạm / Tuyên dương:</div>", unsafe_allow_html=True)
                step_c1, step_c2, step_c3 = st.columns([1, 1, 1])
                with step_c1:
                    if st.button("➖ Giảm", use_container_width=True):
                        if st.session_state.record_times > 1:
                            st.session_state.record_times -= 1
                            st.rerun()
                with step_c2:
                    st.markdown(f"<div style='text-align: center; font-size: 1.4rem; font-weight: 800; padding: 4px;'>{st.session_state.record_times}</div>", unsafe_allow_html=True)
                with step_c3:
                    if st.button("➕ Tăng", use_container_width=True):


